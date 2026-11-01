"""生成简历测试样本：DOCX / TXT，并复用同一份内容保证格式间可对比。

用法：
    python build_resume_samples.py

输出（同目录）：
    sample-resume-java-backend.docx
    sample-resume-java-backend.txt
    resume-template-blank.docx
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent

BODY_FONT = "微软雅黑"
TEXT_COLOR = RGBColor(0x37, 0x41, 0x51)
MUTED_COLOR = RGBColor(0x6B, 0x72, 0x80)


# 用“块”描述简历结构，DOCX 和 TXT 共用一份内容，避免两种格式说的事情不一样。
# 块类型：name 姓名 / tagline 岗位标签 / contact 联系方式 / section 章节 /
#         entry 带日期的条目 / bullet 条目下的小点 / para 普通段落
BLOCKS = [
    ("name", "陆晨"),
    ("tagline", "Java 后端开发工程师 ｜ 3 年经验 ｜ 杭州"),
    ("contact", "手机：138-0000-0000　｜　邮箱：luchen@example.com　｜　GitHub：github.com/luchen-demo"),

    ("section", "教育经历"),
    ("entry", "浙江理工大学　计算机科学与技术（本科）", "2020.09 - 2024.06"),
    ("bullet", "主修课程：数据结构、操作系统、计算机网络、数据库系统原理"),
    ("bullet", "GPA 3.6/4.0，专业排名前 15%，获校级一等奖学金（2022）"),

    ("section", "专业技能"),
    ("bullet", "语言与框架：Java 17/21、Spring Boot 3、Spring MVC、Spring Data JPA、MyBatis"),
    ("bullet", "数据与中间件：MySQL 8、Redis、RabbitMQ、Elasticsearch（了解）"),
    ("bullet", "工程能力：Maven、Git、Docker、Linux、JUnit 5 + Mockito，熟悉常见设计模式"),
    ("bullet", "底层能力：了解 JVM 内存模型与 GC 调优，能用 Arthas 定位线上 CPU 与内存问题"),

    ("section", "工作经历"),
    ("entry", "杭州云帆网络科技有限公司　Java 后端开发工程师", "2024.07 - 至今"),
    ("bullet", "负责智能面试助手平台后端开发，参与需求评审、接口设计与线上问题排查"),
    ("bullet", "主导简历模块从 0 到 1 落地，打通上传、解析、分析与文字面试链路，累计处理简历 2 万+ 份"),
    ("bullet", "统一接口异常与响应结构，收敛重复代码，新增一个接口的平均开发时间缩短约 30%"),

    ("section", "实习经历"),
    ("entry", "杭州星河电商有限公司　后端开发实习生", "2023.09 - 2024.05"),
    ("bullet", "负责订单中心查询接口维护，通过慢 SQL 分析补充联合索引，列表接口 P99 由 1.2s 降低到 260ms"),
    ("bullet", "实现订单状态变更的异步通知，使用 RabbitMQ 削峰，消息积压峰值下降 70%"),
    ("bullet", "编写接口压测脚本，定位连接池配置不合理问题，单机 TPS 提升约 45%"),

    ("section", "项目经历"),
    ("entry", "智能面试助手（Spring Boot 3 + Spring AI + PostgreSQL）　后端负责人", "2024.08 - 至今"),
    ("bullet", "项目描述：面向求职者的模拟面试平台，支持简历上传解析、AI 出题、文字面试与面试报告生成"),
    ("bullet", "负责简历模块整体设计，按 Controller/Service/Repository/Infrastructure 分层，"
               "文件校验、文本解析、对象存储拆成独立组件"),
    ("bullet", "设计基于文件内容哈希的去重方案，重复上传直接返回已有简历，避免对象存储和数据库重复写入"),
    ("bullet", "实现 AI 简历分析的降级策略：结构化输出解析失败先重试一次，仍失败则回退规则评分，保证上传链路不中断"),
    ("bullet", "引入 Redis 缓存简历详情，热点接口 QPS 由 120 提升到 400，数据库读压力下降约 35%"),
    ("entry", "电商订单中心性能优化（Spring Boot + MySQL + Redis）　后端开发", "2023.09 - 2024.05"),
    ("bullet", "负责订单查询链路优化，重构分页查询与索引设计，慢查询数量减少 80%"),
    ("bullet", "设计订单状态机并落地为领域模型，替换 Controller 中的分支判断，线上状态相关缺陷减少 60%"),

    ("section", "自我评价"),
    ("bullet", "习惯先用数据再下结论，优化前后都会保留压测与监控对比，而不是“感觉变快了”"),
    ("bullet", "有文档与注释习惯，接口变更同步更新 README，方便他人接手"),
    ("bullet", "正在系统学习 Spring AI 与 RAG，希望把 AI 能力落到真实业务链路，而不只是跑通示例"),
]

# 空白模板：保留结构，把可替换内容换成占位符，方便改成自己的简历。
TEMPLATE_BLOCKS = [
    ("name", "【姓名】"),
    ("tagline", "【目标岗位】 ｜ 【工作年限】 ｜ 【城市】"),
    ("contact", "手机：【手机号】　｜　邮箱：【邮箱】　｜　GitHub：【仓库地址】"),

    ("section", "求职意向"),
    ("para", "【目标岗位】，期望城市【城市】，可【到岗时间】到岗。"),

    ("section", "教育经历"),
    ("entry", "【学校】　【专业】（【学历】）", "【开始时间】 - 【结束时间】"),
    ("bullet", "主修课程：【课程 1、课程 2、课程 3】"),
    ("bullet", "成绩与荣誉：【GPA / 排名 / 奖学金】"),

    ("section", "专业技能"),
    ("bullet", "语言与框架：【语言、框架、版本】"),
    ("bullet", "数据与中间件：【数据库、缓存、消息队列】"),
    ("bullet", "工程能力：【构建工具、版本管理、容器、测试】"),

    ("section", "工作经历"),
    ("entry", "【公司】　【岗位】", "【开始时间】 - 【结束时间】"),
    ("bullet", "负责【业务范围】，通过【技术手段】解决问题，【可量化结果】"),
    ("bullet", "实现【功能点】，使用【技术方案】，【可量化结果】"),

    ("section", "实习经历"),
    ("entry", "【公司】　【岗位】", "【开始时间】 - 【结束时间】"),
    ("bullet", "负责【业务范围】，【可量化结果】"),

    ("section", "项目经历"),
    ("entry", "【项目名称】（【技术栈】）　【担任角色】", "【开始时间】 - 【结束时间】"),
    ("bullet", "项目描述：【一句话说明业务目标与用户】"),
    ("bullet", "负责【模块/职责】，采用【方案】，解决【具体问题】"),
    ("bullet", "实现【亮点功能】，【可量化结果，如 QPS 提升 40%、耗时降低 60%】"),

    ("section", "自我评价"),
    ("bullet", "【技术偏好或学习习惯】"),
    ("bullet", "【协作方式或文档习惯】"),
]


def set_run_font(run, size=None, bold=None, color=None, font_name=BODY_FONT):
    """设置中英文字体；中文必须额外设置 eastAsia，否则 Word 会回退到宋体。"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_bottom_border(paragraph, color="D1D5DB", size=6):
    """给章节标题加下边框，形成简历常见的分隔线。"""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_section_heading(document, title):
    """章节标题：加粗、字号略大、带下边框。"""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(title)
    set_run_font(run, size=11, bold=True, color=RGBColor(0x11, 0x18, 0x27))
    add_bottom_border(paragraph)


def add_entry(document, left_text, right_text):
    """条目行：左侧公司/项目，右侧时间，用右对齐制表位实现。"""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_TAB_ALIGNMENT.RIGHT)

    left_run = paragraph.add_run(left_text)
    set_run_font(left_run, size=9.5, bold=True, color=RGBColor(0x11, 0x18, 0x27))

    if right_text:
        tab_run = paragraph.add_run("\t" + right_text)
        set_run_font(tab_run, size=8.5, color=MUTED_COLOR)


def add_bullet(document, text):
    """小点：手动加项目符号，避免依赖 numbering，导出 PDF 和复制文本都更稳。"""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08

    bullet_run = paragraph.add_run("• ")
    set_run_font(bullet_run, size=9.5, color=RGBColor(0x25, 0x63, 0xEB))
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=9.5, color=TEXT_COLOR)


def add_paragraph(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5, color=TEXT_COLOR)


def build_document(blocks):
    document = Document()

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(0.9)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal_style = document.styles["Normal"]
    normal_style.font.name = BODY_FONT
    normal_style.font.size = Pt(9.5)
    normal_style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for kind, *values in blocks:
        if kind == "name":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(values[0])
            set_run_font(run, size=19, bold=True, color=RGBColor(0x11, 0x18, 0x27))
        elif kind == "tagline":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(values[0])
            set_run_font(run, size=10.5, bold=True, color=RGBColor(0x25, 0x63, 0xEB))
        elif kind == "contact":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(values[0])
            set_run_font(run, size=9, color=MUTED_COLOR)
        elif kind == "section":
            add_section_heading(document, values[0])
        elif kind == "entry":
            add_entry(document, values[0], values[1] if len(values) > 1 else "")
        elif kind == "bullet":
            add_bullet(document, values[0])
        elif kind == "para":
            add_paragraph(document, values[0])
        else:
            raise ValueError(f"未知块类型: {kind}")

    return document


def to_plain_text(blocks):
    """生成 TXT 版本，保持与 DOCX 相同的段落顺序和标题文字。"""
    lines = []
    for kind, *values in blocks:
        if kind == "section":
            lines.append("")
            lines.append(values[0])
        elif kind == "entry":
            right_text = values[1] if len(values) > 1 else ""
            lines.append(f"{values[0]}    {right_text}".rstrip())
        elif kind == "bullet":
            lines.append(f"- {values[0]}")
        else:
            lines.append(values[0])
    return "\n".join(lines).strip() + "\n"


def rule_based_score(text):
    """复刻 ResumeGradingService 的规则版算法，用于核对兜底路径的预期分值。"""
    sections = sum(1 for keyword in ("教育", "项目", "工作经历", "技能", "实习") if keyword in text)
    skills = sum(1 for keyword in ("Java", "Spring", "Spring Boot", "MySQL", "Redis", "Docker", "Linux", "Git")
                 if keyword in text)

    if len(text) < 200:
        content = 5
    elif len(text) < 500:
        content = 10
    elif len(text) < 1000:
        content = 15
    else:
        content = 18
    content = min(content + sections * 2, 25)

    structure = 6 if sections <= 1 else 10 if sections == 2 else 14 if sections == 3 else 18
    skill_match = 6 if skills <= 1 else 12 if skills <= 3 else 18 if skills <= 5 else 22

    has_action = any(word in text for word in ("负责", "设计", "实现", "优化", "重构"))
    has_metric = any(word in text for word in ("%", "提升", "降低", "减少", "QPS", "耗时"))
    expression = 13 if has_action and has_metric else 10 if has_action else 6

    has_project = "项目" in text
    has_outcome = any(word in text for word in ("提升", "降低", "%", "上线"))
    project = 14 if has_project and has_action and has_outcome else 11 if has_project and has_action else 8 if has_project else 3

    return {
        "content": content,
        "structure": structure,
        "skill": skill_match,
        "expression": expression,
        "project": project,
        "sections": sections,
        "skills": skills,
        "length": len(text),
        "total": content + structure + skill_match + expression + project,
    }


def main():
    sample_docx = BASE_DIR / "sample-resume-java-backend.docx"
    sample_txt = BASE_DIR / "sample-resume-java-backend.txt"
    template_docx = BASE_DIR / "resume-template-blank.docx"

    build_document(BLOCKS).save(sample_docx)
    sample_txt.write_text(to_plain_text(BLOCKS), encoding="utf-8")
    build_document(TEMPLATE_BLOCKS).save(template_docx)

    score = rule_based_score(sample_txt.read_text(encoding="utf-8"))
    print(f"已生成: {sample_docx.name} / {sample_txt.name} / {template_docx.name}")
    print(f"正文长度: {score['length']} 字")
    print(f"章节命中: {score['sections']} 个，技能关键词命中: {score['skills']} 个")
    print(
        "规则版预期分值: "
        f"内容 {score['content']} + 结构 {score['structure']} + 技能 {score['skill']} + "
        f"表达 {score['expression']} + 项目 {score['project']} = {score['total']}"
    )


if __name__ == "__main__":
    main()
